#!/usr/bin/env python3
"""Contact-centric Reverse Boolean Search for LinkedIn Sales Navigator.

Input known buyer/contact names + companies. The engine uses the OpenAI Responses
API (optionally with web search) to:
  1) resolve the exact person/company,
  2) profile management level, company size, and rough revenue,
  3) generate an anonymous Sales Navigator search intended to rediscover the
     known buyer WITHOUT using their name or company,
  4) generate a broader lookalike search for new prospects.

The core `ContactReverseSearchEngine` is standalone and can later be imported
from a FastAPI, Streamlit, or other application.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

DEFAULT_MODEL = os.getenv("OPENAI_MODEL", "gpt-5.6-luna")
MAX_BOOLEAN_OPERATORS_PER_FIELD = 15
DEFAULT_CACHE_DIR = Path(".reverse_boolean_cache")


@dataclass
class KnownContact:
    name: str
    company: str
    location: str | None = None
    website: str | None = None

    def compact(self) -> dict[str, str]:
        out = {
            "name": self.name.strip(),
            "company": self.company.strip(),
        }
        if self.location:
            out["location"] = self.location.strip()
        if self.website:
            out["website"] = self.website.strip()
        return out


SYSTEM_INSTRUCTIONS = r"""
You are the research and search-strategy component of a B2B prospecting tool.
The user gives you KNOWN BUYERS/CONTACTS who already speak with or buy from an
agency. Your task is to reconstruct the kind of buyer and company each person
represents, then generate LinkedIn Sales Navigator search criteria that should
rediscover that known buyer WITHOUT using the person's name or company name in
the actual search criteria.

This is a reverse-search validation workflow:
KNOWN BUYER -> research/profile -> anonymous persona/company profile -> Sales
Navigator search -> human checks whether the known buyer appears -> surrounding
results become candidate lookalikes.

RESEARCH RULES
- Use public web research when available.
- Resolve the exact PERSON and exact COMPANY together. A common name is not
  enough; use company, location, website, job title, and other evidence.
- Never invent a match. If the person/company cannot be resolved confidently,
  mark it unresolved/ambiguous and explain what identifying detail is missing.
- Never claim you viewed private LinkedIn data. You may use public web results,
  public LinkedIn snippets/pages, company websites, business directories,
  credible industry sources, and other public information.
- Distinguish observed facts from estimates.
- If evidence conflicts, report the conflict and lower confidence.
- Every important observed fact or estimate should include a short basis and,
  when available, source URLs in the evidence list.

PERSON PROFILE
For each resolved contact determine:
- current title
- likely function
- management level (Owner/Executive, VP, Director, Manager, Individual
  Contributor, or Unknown)
- Sales Navigator seniority categories that would reasonably capture them
- confidence

COMPANY PROFILE
For each company determine/infer:
- industry / niche
- headquarters/service geography where relevant
- approximate employee count or range
- best matching LinkedIn Company Headcount band
- rough annual revenue range in USD
- confidence and estimation basis

Revenue rules:
- Private-company revenue is often not public. Never fabricate exact revenue.
- Prefer a publicly reported/reputable estimate when available.
- Otherwise provide a BROAD inferred range using employee count, locations,
  service mix, scale, public company signals, industry economics, and other
  available evidence.
- Mark inferred revenue confidence Low or Medium and clearly state the basis.
- If there truly is not enough evidence, use null values rather than fake
  precision.

LINKEDIN SALES NAVIGATOR RULES
- Boolean syntax supports AND, OR, NOT, straight quotes, and parentheses.
- Boolean can be used in the title/company/keyword fields; keep each field to
  about 15 Boolean operators or fewer.
- Generate TITLE Boolean separately from KEYWORDS Boolean.
- Use structured Lead filters separately: Company headcount, Company type,
  Current job title, Function, Seniority level, Geography, Industry, etc.
- Annual Revenue is an ACCOUNT search filter, not a Lead filter. Put revenue
  under account_filters only.
- Do NOT put the known person's name, known company name, company domain, or an
  exact-current-company constraint into the generated search. The point is to
  rediscover the anchor anonymously.
- Search criteria should target peers/lookalikes, not technicians, students,
  recruiters, or irrelevant rank-and-file roles unless the known buyer is
  genuinely at that level.
- Use straight quotes only, never smart quotes.

GENERATE TWO SEARCHES FOR EACH CONTACT
1) validation_search:
   Narrow enough that the known buyer plausibly appears, but anonymous. It may
   use the buyer's actual title family, seniority, niche, company headcount band,
   and geography. It must not use their name/company.
2) expansion_search:
   A modestly broader version for finding new lookalikes. Preserve the buyer
   level and niche/company-scale logic while broadening one or two constraints
   such as geography, related titles, or neighboring headcount bands.

If multiple known contacts are supplied, ALSO create a combined_strategy based
on repeated patterns across the resolved contacts. Do not force dissimilar
contacts into one profile; create multiple archetypes if necessary.

Return ONLY valid JSON with this shape. Use null where unknown:
{
  "summary": "...",
  "contacts": [
    {
      "input": {
        "name": "...",
        "company": "...",
        "location": "... or null",
        "website": "... or null"
      },
      "resolution": {
        "resolved": true,
        "resolved_person": "... or null",
        "resolved_company": "... or null",
        "resolved_company_website": "... or null",
        "confidence": 0.0,
        "ambiguity_note": "... or null"
      },
      "person_profile": {
        "current_title": "... or null",
        "function": "... or null",
        "management_level": "Owner/Executive | VP | Director | Manager | Individual Contributor | Unknown",
        "sales_navigator_seniority": ["..."],
        "confidence": 0.0,
        "basis": "...",
        "evidence": [
          {"claim": "...", "source_url": "... or null"}
        ]
      },
      "company_profile": {
        "industry": "... or null",
        "niche_terms": ["..."],
        "headquarters_or_service_area": "... or null",
        "employee_estimate": {
          "low": 0,
          "high": 0,
          "best_guess": 0,
          "linkedin_headcount_band": "... or null",
          "confidence": "High | Medium | Low | Unknown",
          "basis": "..."
        },
        "revenue_estimate": {
          "low_usd": 0,
          "high_usd": 0,
          "confidence": "High | Medium | Low | Unknown",
          "basis": "...",
          "is_publicly_reported": false
        },
        "evidence": [
          {"claim": "...", "source_url": "... or null"}
        ]
      },
      "validation_search": {
        "title_boolean": "...",
        "keyword_boolean": "...",
        "lead_filters": {
          "company_headcount": ["..."],
          "company_type": ["..."],
          "function": ["..."],
          "seniority": ["..."],
          "geography": ["..."],
          "industry": ["..."]
        },
        "account_filters": {
          "headquarters_location": ["..."],
          "industry": ["..."],
          "company_headcount": ["..."],
          "annual_revenue_min_usd": null,
          "annual_revenue_max_usd": null
        },
        "why_it_should_rediscover_anchor": ["..."],
        "expected_anchor": "Person - Company"
      },
      "expansion_search": {
        "title_boolean": "...",
        "keyword_boolean": "...",
        "lead_filters": {
          "company_headcount": ["..."],
          "company_type": ["..."],
          "function": ["..."],
          "seniority": ["..."],
          "geography": ["..."],
          "industry": ["..."]
        },
        "account_filters": {
          "headquarters_location": ["..."],
          "industry": ["..."],
          "company_headcount": ["..."],
          "annual_revenue_min_usd": null,
          "annual_revenue_max_usd": null
        },
        "what_was_broadened": ["..."],
        "likely_false_positives": ["..."]
      }
    }
  ],
  "combined_strategy": {
    "available": true,
    "shared_buyer_pattern": "...",
    "archetypes": [
      {
        "name": "...",
        "based_on_contacts": ["..."],
        "title_boolean": "...",
        "keyword_boolean": "...",
        "lead_filters": {
          "company_headcount": ["..."],
          "function": ["..."],
          "seniority": ["..."],
          "geography": ["..."],
          "industry": ["..."]
        },
        "why": ["..."]
      }
    ]
  },
  "human_validation_steps": ["..."]
}
""".strip()


def extract_json(text: str) -> dict[str, Any]:
    cleaned = text.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    try:
        value = json.loads(cleaned)
    except json.JSONDecodeError:
        start = cleaned.find("{")
        end = cleaned.rfind("}")
        if start == -1 or end == -1 or end <= start:
            raise ValueError("Model did not return parseable JSON")
        value = json.loads(cleaned[start : end + 1])
    if not isinstance(value, dict):
        raise ValueError("Top-level model output must be a JSON object")
    return value


def count_boolean_operators(query: str | None) -> int:
    if not query:
        return 0
    return len(re.findall(r"\b(?:AND|OR|NOT)\b", query, flags=re.IGNORECASE))


def normalize_quotes(value: str | None) -> str:
    if not value:
        return ""
    return (
        value.replace("“", '"')
        .replace("”", '"')
        .replace("‘", "'")
        .replace("’", "'")
    )


def management_level_from_title(title: str | None) -> str:
    """Deterministic fallback/check for common B2B titles."""
    if not title:
        return "Unknown"
    t = title.casefold()
    owner_exec = (
        "owner", "founder", "co-founder", "cofounder", "president", "chief executive",
        "ceo", "principal", "managing partner", "managing member", "partner",
    )
    if any(term in t for term in owner_exec):
        return "Owner/Executive"
    if "vice president" in t or re.search(r"\bvp\b", t):
        return "VP"
    if "director" in t or "head of" in t:
        return "Director"
    if "manager" in t or "general manager" in t:
        return "Manager"
    return "Individual Contributor"


def linkedin_headcount_band(best_guess: int | None) -> str | None:
    if best_guess is None or best_guess <= 0:
        return None
    bands = [
        (1, 10, "1-10"),
        (11, 50, "11-50"),
        (51, 200, "51-200"),
        (201, 500, "201-500"),
        (501, 1000, "501-1000"),
        (1001, 5000, "1001-5000"),
        (5001, 10000, "5001-10000"),
    ]
    for low, high, label in bands:
        if low <= best_guess <= high:
            return label
    return "10001+"


def format_usd(value: Any) -> str:
    if value is None:
        return "Unknown"
    try:
        n = float(value)
    except (TypeError, ValueError):
        return str(value)
    if n >= 1_000_000_000:
        return f"${n/1_000_000_000:.1f}B"
    if n >= 1_000_000:
        return f"${n/1_000_000:.1f}M"
    if n >= 1_000:
        return f"${n/1_000:.0f}K"
    return f"${n:,.0f}"


def _casefold_text(value: Any) -> str:
    return str(value or "").casefold()


def _search_contains_anchor(search: dict[str, Any], contact: KnownContact) -> list[str]:
    """Check generated query/filter text for accidental anchor leakage."""
    serialized = json.dumps(search, ensure_ascii=False).casefold()
    leaks: list[str] = []
    forbidden = [contact.name.strip(), contact.company.strip()]
    if contact.website:
        domain = re.sub(r"^https?://", "", contact.website.strip(), flags=re.I).split("/")[0]
        forbidden.append(domain)
    for phrase in forbidden:
        phrase = phrase.casefold().strip()
        if len(phrase) >= 4 and phrase in serialized:
            leaks.append(phrase)
    return leaks


def validate_result(result: dict[str, Any], contacts: list[KnownContact]) -> list[str]:
    warnings: list[str] = []
    items = result.get("contacts")
    if not isinstance(items, list) or not items:
        return ["No contact analyses were returned."]

    for idx, item in enumerate(items, start=1):
        if not isinstance(item, dict):
            warnings.append(f"Contact #{idx} result is malformed.")
            continue

        # Compare model's management label with a deterministic title mapping.
        person = item.get("person_profile") or {}
        if isinstance(person, dict):
            title = person.get("current_title")
            model_level = person.get("management_level")
            deterministic = management_level_from_title(title)
            if (
                title
                and model_level
                and deterministic != "Individual Contributor"
                and deterministic != model_level
            ):
                warnings.append(
                    f"Contact #{idx}: title '{title}' maps deterministically to "
                    f"'{deterministic}', while the model returned '{model_level}'. Review manually."
                )

        company = item.get("company_profile") or {}
        if isinstance(company, dict):
            emp = company.get("employee_estimate") or {}
            if isinstance(emp, dict):
                best = emp.get("best_guess")
                if isinstance(best, (int, float)) and best > 0:
                    expected_band = linkedin_headcount_band(int(best))
                    returned_band = emp.get("linkedin_headcount_band")
                    if returned_band and expected_band and returned_band != expected_band:
                        warnings.append(
                            f"Contact #{idx}: employee best guess {best} implies LinkedIn "
                            f"headcount band {expected_band}, but model returned {returned_band}."
                        )

        contact_seed = contacts[idx - 1] if idx - 1 < len(contacts) else None
        for search_name in ("validation_search", "expansion_search"):
            search = item.get(search_name) or {}
            if not isinstance(search, dict):
                continue
            for field in ("title_boolean", "keyword_boolean"):
                query = normalize_quotes(search.get(field))
                search[field] = query
                operators = count_boolean_operators(query)
                if operators > MAX_BOOLEAN_OPERATORS_PER_FIELD:
                    warnings.append(
                        f"Contact #{idx} {search_name}.{field} uses {operators} Boolean "
                        f"operators; target <= {MAX_BOOLEAN_OPERATORS_PER_FIELD}."
                    )
            if contact_seed:
                leaks = _search_contains_anchor(search, contact_seed)
                if leaks:
                    warnings.append(
                        f"Contact #{idx} {search_name} leaks anchor identifiers into search "
                        f"criteria: {', '.join(leaks)}. Remove them before Sales Navigator use."
                    )

    combined = result.get("combined_strategy") or {}
    if isinstance(combined, dict):
        for a_idx, archetype in enumerate(combined.get("archetypes") or [], start=1):
            if not isinstance(archetype, dict):
                continue
            for field in ("title_boolean", "keyword_boolean"):
                query = normalize_quotes(archetype.get(field))
                archetype[field] = query
                operators = count_boolean_operators(query)
                if operators > MAX_BOOLEAN_OPERATORS_PER_FIELD:
                    warnings.append(
                        f"Combined archetype #{a_idx} {field} uses {operators} operators; "
                        f"split it into multiple searches."
                    )
    return warnings


class ContactReverseSearchEngine:
    """Reusable contact-centric engine for integration into the main app later."""

    def __init__(
        self,
        model: str = DEFAULT_MODEL,
        use_web: bool = True,
        search_context_size: str = "low",
        cache_dir: Path | None = DEFAULT_CACHE_DIR,
    ) -> None:
        try:
            from openai import OpenAI
        except ImportError as exc:
            raise RuntimeError("Missing OpenAI SDK. Run: pip install -r requirements.txt") from exc
        self.client = OpenAI()
        self.model = model
        self.use_web = use_web
        self.search_context_size = search_context_size
        self.cache_dir = cache_dir

    def _cache_key(self, contacts: list[KnownContact], agency_context: str) -> str:
        payload = {
            "contacts": [c.compact() for c in contacts],
            "agency_context": agency_context,
            "model": self.model,
            "use_web": self.use_web,
            "search_context_size": self.search_context_size,
            "prompt_version": 2,
        }
        raw = json.dumps(payload, sort_keys=True, ensure_ascii=False).encode("utf-8")
        return hashlib.sha256(raw).hexdigest()

    def _cache_path(self, key: str) -> Path | None:
        if self.cache_dir is None:
            return None
        return self.cache_dir / f"{key}.json"

    def generate(
        self,
        contacts: list[KnownContact],
        agency_context: str = "",
        refresh: bool = False,
    ) -> dict[str, Any]:
        if not contacts:
            raise ValueError("Provide at least 1 known contact.")
        for c in contacts:
            if not c.name.strip() or not c.company.strip():
                raise ValueError("Every contact needs both a person name and company name.")

        key = self._cache_key(contacts, agency_context)
        cache_path = self._cache_path(key)
        if cache_path and cache_path.exists() and not refresh:
            cached = json.loads(cache_path.read_text(encoding="utf-8"))
            cached.setdefault("meta", {})["loaded_from_cache"] = True
            return cached

        payload = {
            "known_contacts": [c.compact() for c in contacts],
            "agency_context": agency_context.strip() or None,
            "objective": (
                "Profile each known buyer and their company, estimate management level, "
                "employee scale and broad revenue, then create anonymous Sales Navigator "
                "validation + expansion searches that can rediscover the buyer and find peers."
            ),
        }

        kwargs: dict[str, Any] = {
            "model": self.model,
            "instructions": SYSTEM_INSTRUCTIONS,
            "input": json.dumps(payload, indent=2),
            "text": {"verbosity": "low"},
        }
        if self.use_web:
            kwargs["tools"] = [
                {
                    "type": "web_search",
                    "search_context_size": self.search_context_size,
                    "user_location": {
                        "type": "approximate",
                        "country": "US",
                        "timezone": "America/New_York",
                    },
                }
            ]
        else:
            # The Responses API does not allow web search and JSON mode in the
            # same request. Without web search, JSON mode gives stricter output;
            # with web search, SYSTEM_INSTRUCTIONS and extract_json() handle it.
            kwargs["text"]["format"] = {"type": "json_object"}

        response = self.client.responses.create(**kwargs)
        result = extract_json(response.output_text)
        result["meta"] = {
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "model": self.model,
            "web_research_enabled": self.use_web,
            "input_contact_count": len(contacts),
            "loaded_from_cache": False,
        }
        result["validation_warnings"] = validate_result(result, contacts)

        if cache_path:
            cache_path.parent.mkdir(parents=True, exist_ok=True)
            cache_path.write_text(json.dumps(result, indent=2), encoding="utf-8")
        return result


# Backwards-friendly alias for easier integration if old code imported the former name.
ReverseBooleanEngine = ContactReverseSearchEngine


def parse_contact_line(line: str) -> KnownContact | None:
    """Accept: Person | Company | optional location | optional website."""
    line = line.strip()
    if not line or line.startswith("#"):
        return None
    parts = [part.strip() for part in line.split("|")]
    if len(parts) < 2 or not parts[0] or not parts[1]:
        raise ValueError(
            "Each line must contain at least: Person Name | Company Name"
        )
    return KnownContact(
        name=parts[0],
        company=parts[1],
        location=parts[2] if len(parts) > 2 and parts[2] else None,
        website=parts[3] if len(parts) > 3 and parts[3] else None,
    )


def load_contacts(path: Path) -> list[KnownContact]:
    if path.suffix.lower() == ".xlsx":
        return load_excel_contacts(path)

    contacts: list[KnownContact] = []
    for line_no, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        try:
            contact = parse_contact_line(line)
        except ValueError as exc:
            raise ValueError(f"{path}:{line_no}: {exc}") from exc
        if contact:
            contacts.append(contact)
    return contacts


def load_excel_contacts(path: Path) -> list[KnownContact]:
    """Load contacts from the first worksheet of an .xlsx workbook."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError(
            "Excel input requires openpyxl. Run: pip install -r requirements.txt"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["Contacts"] if "Contacts" in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if not header_row:
            return []

        def normalize_header(value: Any) -> str:
            return re.sub(r"[^a-z0-9]", "", str(value or "").lower())

        aliases = {
            "name": {"personname", "contactname", "name", "person"},
            "company": {"company", "companyname", "organization", "organisation"},
            "location": {"location", "citystate", "geography"},
            "website": {"website", "companywebsite", "url", "domain"},
        }
        normalized = [normalize_header(value) for value in header_row]
        columns: dict[str, int] = {}
        for field, names in aliases.items():
            match = next((i for i, value in enumerate(normalized) if value in names), None)
            if match is not None:
                columns[field] = match

        missing = [field for field in ("name", "company") if field not in columns]
        if missing:
            raise ValueError(
                f"{path}: Excel header row must include Person Name and Company columns"
            )

        def cell(row: tuple[Any, ...], field: str) -> str | None:
            index = columns.get(field)
            if index is None or index >= len(row) or row[index] is None:
                return None
            value = str(row[index]).strip()
            return value or None

        contacts: list[KnownContact] = []
        for row_no, row in enumerate(rows, start=2):
            name = cell(row, "name")
            company = cell(row, "company")
            if not name and not company:
                continue
            if not name or not company:
                raise ValueError(
                    f"{path}:{row_no}: Person Name and Company are both required"
                )
            contacts.append(
                KnownContact(
                    name=name,
                    company=company,
                    location=cell(row, "location"),
                    website=cell(row, "website"),
                )
            )
        return contacts
    finally:
        workbook.close()


def interactive_contacts() -> list[KnownContact]:
    print("\nEnter known buyers/contacts. A person name AND company are required.")
    print("Location and website are optional but make person/company resolution safer.")
    print("Leave Person Name blank when finished.\n")
    contacts: list[KnownContact] = []
    while True:
        try:
            name = input(f"Contact {len(contacts)+1} - Person Name: ").strip()
        except EOFError:
            break
        if not name:
            break
        company = input("  Company: ").strip()
        if not company:
            print("  ! Company is required. Try this contact again.\n")
            continue
        location = input("  Location (optional): ").strip() or None
        website = input("  Company website (optional): ").strip() or None
        contacts.append(KnownContact(name, company, location, website))
        print()
    return contacts


def print_filters(filters: Any, indent: str = "  ") -> None:
    if not isinstance(filters, dict):
        return
    for key, values in filters.items():
        if values in (None, [], ""):
            continue
        if isinstance(values, list):
            text = ", ".join(map(str, values))
        elif key.endswith("_usd"):
            text = format_usd(values)
        else:
            text = str(values)
        print(f"{indent}{key.replace('_', ' ').title()}: {text}")


def print_search(search: dict[str, Any], label: str) -> None:
    print(f"\n{label}")
    print("-" * len(label))
    title_q = normalize_quotes(search.get("title_boolean"))
    keyword_q = normalize_quotes(search.get("keyword_boolean"))
    print("\nTITLE BOOLEAN (paste into Current job title/title field):")
    print(title_q or "[none]")
    print(f"[Boolean operators: {count_boolean_operators(title_q)}]")
    print("\nKEYWORD BOOLEAN:")
    print(keyword_q or "[none]")
    print(f"[Boolean operators: {count_boolean_operators(keyword_q)}]")
    print("\nLEAD FILTERS:")
    print_filters(search.get("lead_filters"))
    print("\nACCOUNT FILTERS (optional pre-qualification):")
    print_filters(search.get("account_filters"))


def print_strategy(result: dict[str, Any]) -> None:
    print("\n" + "=" * 84)
    print("CONTACT-CENTRIC REVERSE SALES NAVIGATOR SEARCH")
    print("=" * 84)
    if result.get("summary"):
        print(f"\n{result['summary']}")

    for idx, item in enumerate(result.get("contacts") or [], start=1):
        if not isinstance(item, dict):
            continue
        inp = item.get("input") or {}
        resolution = item.get("resolution") or {}
        person = item.get("person_profile") or {}
        company = item.get("company_profile") or {}

        print("\n" + "#" * 84)
        print(f"ANCHOR {idx}: {inp.get('name', 'Unknown')} @ {inp.get('company', 'Unknown')}")
        print("#" * 84)
        print(
            f"Resolved: {'YES' if resolution.get('resolved') else 'NO'} | "
            f"Confidence: {resolution.get('confidence', 'Unknown')}"
        )
        if resolution.get("ambiguity_note"):
            print(f"Resolution note: {resolution['ambiguity_note']}")

        print("\nBUYER PROFILE")
        print(f"  Current title: {person.get('current_title') or 'Unknown'}")
        print(f"  Management level: {person.get('management_level') or 'Unknown'}")
        print(f"  Function: {person.get('function') or 'Unknown'}")
        seniority = person.get("sales_navigator_seniority") or []
        print(f"  Sales Nav seniority: {', '.join(map(str, seniority)) if seniority else 'Unknown'}")

        print("\nCOMPANY PROFILE")
        print(f"  Industry: {company.get('industry') or 'Unknown'}")
        emp = company.get("employee_estimate") or {}
        if isinstance(emp, dict):
            low, high, best = emp.get("low"), emp.get("high"), emp.get("best_guess")
            if best:
                print(f"  Employees: ~{best} (range {low or '?'}-{high or '?'})")
            else:
                print("  Employees: Unknown")
            print(f"  LinkedIn headcount band: {emp.get('linkedin_headcount_band') or 'Unknown'}")
            print(f"  Employee confidence: {emp.get('confidence') or 'Unknown'}")
        rev = company.get("revenue_estimate") or {}
        if isinstance(rev, dict):
            low, high = rev.get("low_usd"), rev.get("high_usd")
            if low is not None or high is not None:
                print(f"  Rough annual revenue: {format_usd(low)} - {format_usd(high)}")
            else:
                print("  Rough annual revenue: Unknown")
            print(f"  Revenue confidence: {rev.get('confidence') or 'Unknown'}")
            if rev.get("basis"):
                print(f"  Revenue basis: {rev['basis']}")

        validation = item.get("validation_search") or {}
        expansion = item.get("expansion_search") or {}
        if isinstance(validation, dict):
            print_search(validation, "VALIDATION SEARCH — should rediscover this known buyer")
            why = validation.get("why_it_should_rediscover_anchor") or []
            if why:
                print("\nWHY THE ANCHOR SHOULD MATCH:")
                for reason in why:
                    print(f"  - {reason}")
        if isinstance(expansion, dict):
            print_search(expansion, "EXPANSION SEARCH — use surrounding results as lookalike prospects")
            broadened = expansion.get("what_was_broadened") or []
            if broadened:
                print("\nBROADENED FROM VALIDATION:")
                for reason in broadened:
                    print(f"  - {reason}")

    combined = result.get("combined_strategy") or {}
    archetypes = combined.get("archetypes") if isinstance(combined, dict) else None
    if archetypes:
        print("\n" + "=" * 84)
        print("COMBINED / REPEATED BUYER PATTERNS")
        print("=" * 84)
        if combined.get("shared_buyer_pattern"):
            print(f"\n{combined['shared_buyer_pattern']}")
        for idx, archetype in enumerate(archetypes, start=1):
            if not isinstance(archetype, dict):
                continue
            print(f"\nARCHETYPE {idx}: {archetype.get('name', 'Unnamed')}")
            print(f"Based on: {', '.join(map(str, archetype.get('based_on_contacts') or []))}")
            print(f"TITLE: {normalize_quotes(archetype.get('title_boolean'))}")
            print(f"KEYWORDS: {normalize_quotes(archetype.get('keyword_boolean'))}")
            print("FILTERS:")
            print_filters(archetype.get("lead_filters"))

    warnings = result.get("validation_warnings") or []
    if warnings:
        print("\n" + "=" * 84)
        print("PROGRAM VALIDATION WARNINGS")
        print("=" * 84)
        for warning in warnings:
            print(f"  ! {warning}")

    steps = result.get("human_validation_steps") or []
    if steps:
        print("\nHUMAN VALIDATION")
        for idx, step in enumerate(steps, start=1):
            print(f"  {idx}. {step}")


def save_result(result: dict[str, Any], output_path: Path) -> None:
    output_path.write_text(json.dumps(result, indent=2), encoding="utf-8")


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Profile known buyers and reverse-engineer anonymous LinkedIn Sales Navigator "
            "searches that should rediscover them and surface lookalikes."
        )
    )
    parser.add_argument(
        "--contacts",
        type=Path,
        help="Text file: Person Name | Company | optional location | optional website",
    )
    parser.add_argument(
        "--agency-context",
        default="",
        help='Optional context, e.g. "Agency sells AI lead follow-up to HVAC contractors".',
    )
    parser.add_argument(
        "--model",
        default=DEFAULT_MODEL,
        help=f"OpenAI model (default: {DEFAULT_MODEL}; can also set OPENAI_MODEL).",
    )
    parser.add_argument(
        "--no-web",
        action="store_true",
        help="Disable API web research. Not recommended with only person/company names.",
    )
    parser.add_argument(
        "--search-context",
        choices=["low", "medium", "high"],
        default="low",
        help="Web-search context size (default: low for MVP cost control).",
    )
    parser.add_argument(
        "--refresh",
        action="store_true",
        help="Ignore cached result and pay for fresh research.",
    )
    parser.add_argument(
        "--no-cache",
        action="store_true",
        help="Disable local response cache.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("contact_boolean_strategy.json"),
        help="JSON output path (default: contact_boolean_strategy.json).",
    )
    return parser


def main(argv: Iterable[str] | None = None) -> int:
    try:
        from dotenv import load_dotenv
    except ImportError:
        load_dotenv = None
    if load_dotenv is not None:
        load_dotenv()

    args = build_parser().parse_args(argv)

    if not os.getenv("OPENAI_API_KEY"):
        print("ERROR: OPENAI_API_KEY is not set.", file=sys.stderr)
        print("Copy .env.example to .env and add your API key.", file=sys.stderr)
        return 2

    try:
        contacts = load_contacts(args.contacts) if args.contacts else interactive_contacts()
    except (OSError, ValueError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2

    if not contacts:
        print("ERROR: provide at least 1 known contact.", file=sys.stderr)
        return 2

    print(f"\nAnalyzing {len(contacts)} known contact(s) with {args.model}...")
    print(f"Web research: {'OFF' if args.no_web else 'ON'}")
    print(f"Cache: {'OFF' if args.no_cache else 'ON'}")

    engine = ContactReverseSearchEngine(
        model=args.model,
        use_web=not args.no_web,
        search_context_size=args.search_context,
        cache_dir=None if args.no_cache else DEFAULT_CACHE_DIR,
    )

    try:
        result = engine.generate(
            contacts,
            agency_context=args.agency_context,
            refresh=args.refresh,
        )
    except Exception as exc:
        print(f"\nERROR: {exc}", file=sys.stderr)
        return 1

    print_strategy(result)
    save_result(result, args.output)
    print(f"\nFull structured result saved to: {args.output.resolve()}")
    if result.get("meta", {}).get("loaded_from_cache"):
        print("Result came from local cache (no new API research call). Use --refresh to rerun.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
